"""Writes Receipt objects out to a formatted .xlsx file."""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HEADERS = ["Date", "Receipt Number", "Amount"]


def export_receipts(receipts, output_path):
    """Write receipts to an .xlsx file at output_path, sorted by date."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "GrabExpress Receipts"

    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for receipt in sorted(receipts, key=lambda r: r.date):
        row = [
            datetime.strptime(receipt.date, "%Y-%m-%d"),
            receipt.receipt_no,
            receipt.amount,
        ]
        sheet.append(row)

    date_column = 1
    amount_column = 3
    for row in sheet.iter_rows(min_row=2, min_col=date_column, max_col=date_column):
        row[0].number_format = "YYYY-MM-DD"
    for row in sheet.iter_rows(min_row=2, min_col=amount_column, max_col=amount_column):
        row[0].number_format = "0.00"

    _auto_size_columns(sheet)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def _auto_size_columns(sheet):
    for col_idx, column_cells in enumerate(sheet.columns, start=1):
        max_length = max(len(str(cell.value)) for cell in column_cells if cell.value is not None)
        sheet.column_dimensions[get_column_letter(col_idx)].width = max_length + 4
